import json,re,hashlib,zipfile
from pathlib import Path
from collections import Counter
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import pypdfium2 as pdfium

ROOT=Path(__file__).resolve().parent
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def norm(s):return re.sub(r'\s+','',s)
def main():
    info=json.loads((ROOT/'build_info.json').read_text(encoding='utf-8'))
    qs=json.loads((ROOT/'questions.json').read_text(encoding='utf-8'))
    docpath=Path(info['word']);pdfpath=ROOT/'preview.pdf'
    assert sha(docpath)==info['word_sha256']
    assert pdfpath.stat().st_mtime>=docpath.stat().st_mtime
    doc=Document(docpath);ps=doc.paragraphs
    qid=info['qid'];aid=info['aid']
    def nid(p):
        el=p._p.find('w:pPr/w:numPr/w:numId',{'w':qn('w:p')[1:].split('}')[0]})
        return int(el.get(qn('w:val'))) if el is not None else None
    assert sum(nid(p)==qid for p in ps)==110
    assert sum(nid(p)==aid for p in ps)==110
    for q in qs:
        n=q['n'];p=ps[info['question_paragraphs'][str(n)]]
        assert p.text==q['stem'] and nid(p)==qid
        p=ps[info['analysis_paragraphs'][str(n)]]
        assert p.text==f'答案：{q["answer"]}　{q["point"]}' and nid(p)==aid
        i=info['analysis_paragraphs'][str(n)]
        assert [x.text for x in ps[i+1:i+1+len(q['analysis'])]]==q['analysis']
    assert len(doc.tables)==4
    tablepairs=[]
    for t in doc.tables:
        rows=len(t.rows)-1
        for pair in range(5):
            for r in range(1,rows+1):
                n=t.cell(r,2*pair).text;a=t.cell(r,2*pair+1).text
                if n:tablepairs.append((int(n),a))
    assert tablepairs==[(q['n'],q['answer']) for q in qs]
    option_num_counts=Counter(nid(p) for p in ps if p.style.name=='Option')
    assert len(option_num_counts)==110 and all(v==4 for v in option_num_counts.values())
    with zipfile.ZipFile(docpath) as z:
        assert z.testzip() is None
        styles=etree.fromstring(z.read('word/styles.xml'))
        assert not list(styles.iter(qn('w:pBdr')))
        embedded=[x for x in z.namelist() if x.startswith('word/media/')]
        assert len(embedded)==2
        assert {hashlib.sha256(z.read(x)).hexdigest() for x in embedded}==set(info['source_images'].values())
    textpages=[p.extract_text() for p in PdfReader(pdfpath).pages]
    assert all(len(norm(t))>15 for t in textpages),'blank page'
    answer_start=next(i for i,t in enumerate(textpages) if '参考答案与解析' in norm(t))
    student_pages=textpages[:answer_start]
    q_pages={};a_pages={};matches={}
    for q in qs:
        n=q['n']
        hits=[i for i,t in enumerate(student_pages) if re.search(rf'(?m)^\s*{n}\.\s',t)]
        assert len(hits)==1,(n,hits)
        i=hits[0];q_pages[n]=i+1
        wanted=[q['stem'],q['ask']]+q['options']
        assert all(norm(s) in norm(textpages[i]) for s in wanted),(n,'question/option split, missing or mismatch')
        ahits=[j for j in range(answer_start,len(textpages)) if re.search(rf'(?m)^\s*{n}\.\s*答案',textpages[j])]
        assert len(ahits)==1,(n,ahits)
        j=ahits[0];a_pages[n]=j+1
        assert norm(f'答案：{q["answer"]}') in norm(textpages[j])
        assert all(norm(s) in norm(textpages[j]) for s in q['analysis']),(n,'analysis missing/split')
        matches[n]=dict(question_page=i+1,analysis_page=j+1,text_and_options_present=True,analysis_paragraphs_present=True)
    groups=json.loads((ROOT/'groups_formal.json').read_text(encoding='utf-8'))
    for n,material in groups.items():
        p=textpages[q_pages[int(n)]-1]
        assert all(norm(x) in norm(p) for x in material),(n,'material not with first item')
    render=ROOT/'render';render.mkdir(exist_ok=True)
    for old in render.glob('page-*.png'):old.unlink()
    pdf=pdfium.PdfDocument(str(pdfpath));pages=[]
    for i,page in enumerate(pdf):
        dest=render/f'page-{i+1:02}.png';im=page.render(scale=150/72).to_pil();im.save(dest)
        pages.append(dict(page=i+1,path=str(dest),sha256=sha(dest),width=im.width,height=im.height,content_chars=len(norm(textpages[i]))))
    (ROOT/'page_text.json').write_text(json.dumps({str(i+1):t for i,t in enumerate(textpages)},ensure_ascii=False,indent=2),encoding='utf-8')
    manifest=dict(word_sha256=sha(docpath),pdf_sha256=sha(pdfpath),dpi=150,pages=pages)
    (ROOT/'page_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    report=dict(passed=True,word_sha256=sha(docpath),pdf_sha256=sha(pdfpath),pages=len(textpages),
      answer_table_start_page=answer_start+1,questions=110,analyses=110,native_question_numbers=110,native_analysis_numbers=110,
      native_option_groups=110,embedded_images=2,answer_table_matches=True,no_question_block_split=True,no_analysis_block_split=True,
      shared_material_with_first_question=True,items=matches,
      scope='OOXML、题目/四项/全部解析文字绑定、原生编号、答案表、嵌图指纹及PDF单页完整性；视觉质量另由逐页实际查看确认。')
    (ROOT/'document_verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:report[k] for k in ['passed','pages','answer_table_start_page','questions','analyses','native_option_groups']},ensure_ascii=False))
if __name__=='__main__':main()
