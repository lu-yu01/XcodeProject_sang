"""One data source, a new 110-question Word practice book, no prior-file overwrite."""
import json,hashlib,importlib,math
from pathlib import Path
from collections import Counter
from docx import Document
from docx.shared import Pt,Mm,Inches,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from layout import style,Numbers,keep,picture,page_field

ROOT=Path(__file__).resolve().parent
BASE=ROOT.parent.parent
NAME='国考判断推理进阶专项练习-110题（含答案解析）'
OUT=BASE/'模式题'/(NAME+'.docx')
INDEX=BASE/'知识库'/'套卷索引'/(NAME+'.md')
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def table(doc,qs):
    rows=math.ceil(len(qs)/5)
    t=doc.add_table(rows=rows+1,cols=10);t.autofit=False
    pr=t._tbl.tblPr
    width=pr.find(qn('w:tblW'));width.set(qn('w:type'),'dxa');width.set(qn('w:w'),'9630')
    border=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        border.append(e)
    pr.append(border)
    mar=OxmlElement('w:tblCellMar')
    for side,value in [('top',65),('bottom',65),('start',75),('end',75)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');mar.append(e)
    pr.append(mar)
    for r,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c,cell in enumerate(row.cells):
            cell.width=Inches(.66875)
            tcw=cell._tc.get_or_add_tcPr().find(qn('w:tcW'));tcw.set(qn('w:w'),'963');tcw.set(qn('w:type'),'dxa')
            i=(c//2)*rows+(r-1)
            txt=('题号' if c%2==0 else '答案') if r==0 else (('' if i>=len(qs) else str(qs[i]['n'])) if c%2==0 else ('' if i>=len(qs) else qs[i]['answer']))
            p=cell.paragraphs[0];p.style='AnswerTable';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(txt).bold=r==0 or c%2==1
            va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');cell._tc.get_or_add_tcPr().append(va)
            if r==0:
                e=OxmlElement('w:shd');e.set(qn('w:fill'),'EAEAEA');cell._tc.get_or_add_tcPr().append(e)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return t

def main():
    modules=[importlib.import_module(x) for x in ['analogy','definitions','formal','arguments']]
    reports={m.__name__:m.verify() for m in modules}
    modules[0].draw_all()
    qs=sum([m.QUESTIONS for m in modules],[])
    assert [q['n'] for q in qs]==list(range(1,111))
    assert Counter(q['category'] for q in qs)=={'类比推理':20,'定义判断':30,'逻辑推理':30,'逻辑论证':30}
    assert all(len(q['options'])==len(set(q['options']))==4 and q['answer'] in 'ABCD' for q in qs)
    assert all(isinstance(q['analysis'],list) and q['analysis'] and all(isinstance(x,str) and x for x in q['analysis']) for q in qs)
    assert all(v.get('passed',v.get('passed_structural_and_explicit_checks',False)) for v in reports.values())
    groups=modules[2].GROUPS
    assert set(groups)=={51,56,61}
    def context(q):
        return next((''.join(groups[start]) for start in groups if start<=q['n']<=start+4),'')
    assert len({context(q)+q['stem']+q['ask']+'|'.join(q['options'])+str(q.get('image')) for q in qs})==110
    (ROOT/'questions.json').write_text(json.dumps(qs,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'model_verification.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2),encoding='utf-8')
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=Mm(18);sec.bottom_margin=Mm(18);sec.left_margin=Mm(20);sec.right_margin=Mm(20)
    sec.header_distance=Mm(8);sec.footer_distance=Mm(8)
    # python-docx template contains a coloured title border unrelated to this book.
    for st in doc.styles:
        for e in list(st.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
    style(doc.styles['Normal'],11,after=4,line=1.16)
    for name,size,bold,before,after in [('Title',22,True,0,8),('Subtitle',11,False,0,9),('Heading 1',16,True,9,8),('Heading 2',12,True,7,5),('Heading 3',11,True,5,4)]:
        style(doc.styles[name],size,bold,before,after,line=1.12,east='黑体' if bold else '宋体')
    for name,size,after in [('Question',11,4),('Option',11,3),('Ask',11,4),('Material',11,4),('Analysis',10.5,4),('AnswerTable',9,2),('Figure',10,3),('Small',9,4)]:
        if name not in doc.styles:doc.styles.add_style(name,1)
        style(doc.styles[name],size,after=after,line=1.14)
    style(doc.styles['Header'],8,after=0,line=1)
    style(doc.styles['Footer'],8,after=0,line=1)
    doc.styles['Heading 1'].paragraph_format.keep_with_next=True
    doc.styles['Heading 2'].paragraph_format.keep_with_next=True
    header=sec.header.paragraphs[0];header.text='国考判断推理进阶专项练习 · 110题';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;page_field(foot)
    doc.core_properties.title='国考判断推理进阶专项练习（110题）'
    doc.core_properties.subject='类比20 / 定义30 / 逻辑推理30 / 逻辑论证30'
    doc.core_properties.author='判断推理命题与质检工程'
    title=doc.add_paragraph('国考判断推理进阶专项练习',style='Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('原创题目册 110题',style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('编排说明',style='Heading 2')
    for s in [
      '本册为模块专项练习：类比推理20题、定义判断30题、逻辑推理30题、逻辑论证30题。可按模块分次完成，所有题目均为单项选择。',
      '逻辑推理中，第51—55、56—60、61—65题分别共用一组材料；各题新增条件仅对本题有效。第66—80题为独立题。',
      '题型与信息结构主要参考2022—2026年国考，地方考试辅助拓展；本册配额为专项训练设置。题中场景、机构和数字为练习设定，按所给条件作答。',
      '题目集中在前，答案表与逐题解析集中在后。定义判断以题干给出的定义为准；类比关系图按概念间关系理解。']:
        doc.add_paragraph(s)
    doc.add_paragraph('一 类比推理 第1至20题',style='Heading 1')
    nums=Numbers(doc);qid=nums.make();aid=nums.make()
    starts={21:'二 定义判断 第21至50题',51:'三 逻辑推理 第51至80题',81:'四 逻辑论证 第81至110题'}
    qparagraphs={};aparagraphs={}
    for q in qs:
        n=q['n']
        if n in starts:
            p=doc.add_paragraph(starts[n],style='Heading 1');p.paragraph_format.page_break_before=True
        if n in groups:
            p=doc.add_paragraph(f'第{n}至{n+4}题共用材料',style='Heading 2')
            if n!=51:p.paragraph_format.page_break_before=True
            material=[p]
            for text in groups[n]:material.append(doc.add_paragraph(text,style='Material'))
            keep(material)
            material[-1].paragraph_format.keep_with_next=True
        if n==66:
            p=doc.add_paragraph('独立题 第66至80题',style='Heading 2');p.paragraph_format.page_break_before=True
        paras=[]
        p=doc.add_paragraph(q['stem'],style='Question');nums.apply(p,qid);paras.append(p);qparagraphs[n]=len(doc.paragraphs)-1
        if q.get('image'):
            paras.append(picture(doc,ROOT/q['image'],f'第{n}题概念关系图',max_width=3.7,max_height=2.59))
        if q['ask']:paras.append(doc.add_paragraph(q['ask'],style='Ask'))
        oid=nums.make('upperLetter','%1.',300,300)
        for text in q['options']:
            p=doc.add_paragraph(text,style='Option');nums.apply(p,oid);paras.append(p)
        keep(paras)
        paras[-1].paragraph_format.space_after=Pt(9)
    p=doc.add_paragraph('参考答案与解析',style='Heading 1');p.paragraph_format.page_break_before=True
    for start,end,label in [(1,20,'类比推理'),(21,50,'定义判断'),(51,80,'逻辑推理'),(81,110,'逻辑论证')]:
        doc.add_paragraph(f'{label}答案',style='Heading 2');table(doc,qs[start-1:end])
    for q in qs:
        n=q['n']
        if n in [1,21,51,81]:
            p=doc.add_paragraph({1:'类比推理解析',21:'定义判断解析',51:'逻辑推理解析',81:'逻辑论证解析'}[n],style='Heading 1');p.paragraph_format.page_break_before=True
        ps=[]
        p=doc.add_paragraph(style='Heading 3');nums.apply(p,aid);p.add_run(f'答案：{q["answer"]}　{q["point"]}');ps.append(p);aparagraphs[n]=len(doc.paragraphs)-1
        for s in q['analysis']:ps.append(doc.add_paragraph(s,style='Analysis'))
        keep(ps)
        ps[-1].paragraph_format.space_after=Pt(8)
    OUT.parent.mkdir(exist_ok=True);doc.save(OUT)
    info=dict(word=str(OUT),word_sha256=sha(OUT),questions=110,category_counts=dict(Counter(q['category'] for q in qs)),
      source_sha256=sha(ROOT/'questions.json'),answer_counts=dict(Counter(q['answer'] for q in qs)),
      qid=qid,aid=aid,question_paragraphs=qparagraphs,analysis_paragraphs=aparagraphs,
      native_list_numbering=True,source_images={q['image']:sha(ROOT/q['image']) for q in qs if q.get('image')},
      status='built; independent review and final render QA pending')
    (ROOT/'build_info.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
    if not INDEX.exists():
        INDEX.write_text(f'# {NAME}\n\n- Word：[成品](../../模式题/{NAME}.docx)\n- 工作目录：[源稿与检查记录](../../.work/{ROOT.name}/)\n- 结构：类比20、定义30、逻辑推理30、逻辑论证30；逻辑推理含3组一拖五15题与15道独立题。\n- 状态：已构建，独立试解、最终渲染及逐页检查进行中，尚未标记可交付。\n- Markdown仅为索引与检查记录，完整题目见Word。\n',encoding='utf-8')
    print(json.dumps({k:info[k] for k in ['word','word_sha256','questions','category_counts','answer_counts']},ensure_ascii=False))
if __name__=='__main__':main()
