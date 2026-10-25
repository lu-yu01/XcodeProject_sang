"""Generate the requested 55-item exercise and its knowledge-base index."""
from pathlib import Path
from collections import Counter
import json
import sys
from docx import Document
from docx.shared import Mm
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import drawkit as k
from content import TITLE,SUBTITLE,INTRO,QUESTIONS,GROUPS,draw_all
from layout import style,Numbers,keep,page_field,picture,answer_table
from suite_record import DOCX,write_record,ROOT

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    draw_all()
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(20)
    sec.header_distance=Mm(8);sec.footer_distance=Mm(9)
    doc.core_properties.title=TITLE;doc.core_properties.subject=SUBTITLE;doc.core_properties.author='教研资料'
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    style(doc.styles['Normal'])
    style(doc.styles['Title'],20,True,0,6,1.1,'黑体')
    style(doc.styles['Subtitle'],10.5,False,0,5,1.15)
    style(doc.styles['Heading 1'],15,True,12,8,1.1,'黑体');doc.styles['Heading 1'].paragraph_format.keep_with_next=True
    style(doc.styles['Heading 2'],12,True,10,6,1.1,'黑体');doc.styles['Heading 2'].paragraph_format.keep_with_next=True
    params={
        'Question':dict(size=11,before=8,after=4),
        'Ask':dict(size=11,before=0,after=3),
        'Option':dict(size=11,before=0,after=2),
        'Figure':dict(size=1,before=2,after=4,line=1),
        'Common':dict(size=10.5,before=0,after=3,line=1.15),
        'AnswerTitle':dict(size=11,bold=True,before=7,after=3,east='黑体'),
        'Point':dict(size=10,bold=True,before=0,after=2),
        'Analysis':dict(size=10.5,before=0,after=3,line=1.15),
        'Small':dict(size=9.5,before=0,after=5,line=1.15),
        'AnswerTable':dict(size=10,before=0,after=0,line=1.05),
    }
    for name,values in params.items():
        st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);st.base_style=doc.styles['Normal'];style(st,**values)
    for name in ('AnswerTitle','Point'):doc.styles[name].paragraph_format.keep_with_next=True
    for name in ('Header','Footer'):style(doc.styles[name],8,False,0,0,1)
    sec.header.paragraphs[0].text=TITLE
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('第 ');page_field(footer);footer.add_run(' 页')
    for text_,s in [(TITLE,'Title'),(SUBTITLE,'Subtitle')]:
        p=doc.add_paragraph(text_,s);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(INTRO,'Small')
    doc.add_paragraph('姓名：________________　日期：________________　得分：________','Small')
    nums=Numbers(doc);question_id=nums.make();answer_id=nums.make()
    headings={1:'一 图形推理 共30题',31:'二 逻辑判断一拖五 共15题',46:'三 逻辑判断加强削弱 共10题'}
    for q in QUESTIONS:
        n=q['n']
        if n in headings:
            if n!=1:doc.add_page_break()
            doc.add_paragraph(headings[n],'Heading 1')
        if n in GROUPS:
            if n!=31:doc.add_page_break()
            material=[]
            for i,s in enumerate(GROUPS[n]):
                p=doc.add_paragraph(s,'Common')
                if i==0:p.runs[0].bold=True
                material.append(p)
            for p in material:p.paragraph_format.keep_with_next=True
        block=[]
        p=doc.add_paragraph(q['stem'],'Question');nums.apply(p,question_id);block.append(p)
        if q['ask']:block.append(doc.add_paragraph(q['ask'],'Ask'))
        if q['image']:block.append(picture(doc,k.ASSETS/q['image'],f"第{n}题题图和图形选项"))
        if q['options']:
            oid=nums.make('upperLetter',left=450,hang=300)
            for option in q['options']:
                p=doc.add_paragraph(option,'Option');nums.apply(p,oid);block.append(p)
        keep(block)
    doc.add_page_break()
    p=doc.add_paragraph('参考答案与解析','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('答案速查','Heading 2');answer_table(doc,QUESTIONS)
    for q in QUESTIONS:
        block=[]
        p=doc.add_paragraph('答案：'+q['answer'],'AnswerTitle');nums.apply(p,answer_id);block.append(p)
        block.append(doc.add_paragraph('考点：'+q['point'],'Point'))
        for i,s in enumerate(q['analysis']):block.append(doc.add_paragraph(('解析：' if i==0 else '')+s,'Analysis'))
        if q['explanation_image']:block.append(picture(doc,k.ASSETS/q['explanation_image'],f"第{q['n']}题解析辅助图",6.2,2.5))
        keep(block)
    DOCX.parent.mkdir(exist_ok=True);doc.save(DOCX)
    meta={'word':str(DOCX),'questions':55,'categories':dict(Counter(q['category'] for q in QUESTIONS)),
          'answers':dict(Counter(q['answer'] for q in QUESTIONS)),'inline_images':len(doc.inline_shapes),
          'markdown':str(write_record())}
    (ROOT/'build_info.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'figure_layout.json').write_text(json.dumps(k.LAYOUTS,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(meta,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
