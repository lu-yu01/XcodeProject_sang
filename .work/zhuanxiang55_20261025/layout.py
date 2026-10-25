"""A4 Chinese exercise layout and true Word list numbering."""
from docx.shared import Pt,RGBColor,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

def style(st,size=11,bold=False,before=0,after=4,line=1.2,east='宋体'):
    st.font.name='Calibri';st.font.size=Pt(size);st.font.bold=bold;st.font.italic=False;st.font.color.rgb=RGBColor(0,0,0)
    rf=st.element.get_or_add_rPr().get_or_add_rFonts()
    for key,value in [('ascii','Calibri'),('hAnsi','Calibri'),('eastAsia',east)]:rf.set(qn('w:'+key),value)
    pf=st.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=line;pf.widow_control=True
    sg=OxmlElement('w:snapToGrid');sg.set(qn('w:val'),'0');st.element.get_or_add_pPr().append(sg)

class Numbers:
    def __init__(self,doc):
        self.root=doc.part.numbering_part.element
        self.a=max([int(e.get(qn('w:abstractNumId'))) for e in self.root if e.tag==qn('w:abstractNum')]+[0])+1
        self.n=max([int(e.get(qn('w:numId'))) for e in self.root if e.tag==qn('w:num')]+[0])+1
    def make(self,fmt='decimal',label='%1.',left=300,hang=300):
        aid,nid=self.a,self.n;self.a+=1;self.n+=1
        a=OxmlElement('w:abstractNum');a.set(qn('w:abstractNumId'),str(aid))
        nsid=OxmlElement('w:nsid');nsid.set(qn('w:val'),f'{0xA6000000+aid:08X}');a.append(nsid)
        multi=OxmlElement('w:multiLevelType');multi.set(qn('w:val'),'singleLevel');a.append(multi)
        lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
        for tag,val in [('start','1'),('numFmt',fmt),('suff','space'),('lvlText',label),('lvlJc','left')]:
            e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);lvl.append(e)
        pp=OxmlElement('w:pPr');ind=OxmlElement('w:ind');ind.set(qn('w:left'),str(left));ind.set(qn('w:hanging'),str(hang));pp.append(ind);lvl.append(pp);a.append(lvl)
        first=next((e for e in self.root if e.tag==qn('w:num')),None)
        if first is None:self.root.append(a)
        else:first.addprevious(a)
        n=OxmlElement('w:num');n.set(qn('w:numId'),str(nid));x=OxmlElement('w:abstractNumId');x.set(qn('w:val'),str(aid));n.append(x);self.root.append(n)
        return nid
    def apply(self,p,nid):
        np=OxmlElement('w:numPr')
        for tag,val in [('ilvl','0'),('numId',str(nid))]:
            e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);np.append(e)
        p._p.get_or_add_pPr().append(np)

def keep(paras):
    for i,p in enumerate(paras):
        p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=i<len(paras)-1

def page_field(p):
    run=p.add_run()
    for typ,txt in [('begin',None),(None,' PAGE '),('end',None)]:
        if typ:e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),typ)
        else:e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=txt
        run._r.append(e)

def picture(doc,path,description,max_width=6.6,max_height=4.1):
    p=doc.add_paragraph(style='Figure');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    with Image.open(path) as im:w,h=im.size
    shape=p.add_run().add_picture(str(path),width=Inches(min(max_width,max_height*w/h)))
    shape._inline.docPr.set('descr',description)
    return p

def answer_table(doc,questions):
    t=doc.add_table(rows=12,cols=10);t.autofit=False
    pr=t._tbl.tblPr;tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'9630')
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side)
        for name,value in [('val','single'),('sz','5'),('color','D9D9D9')]:e.set(qn('w:'+name),value)
        borders.append(e)
    pr.append(borders)
    mar=OxmlElement('w:tblCellMar')
    for side,value in [('top',80),('bottom',80),('start',80),('end',80)]:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');mar.append(e)
    pr.append(mar)
    for col in t.columns:col.width=Inches(963/1440)
    for r,row in enumerate(t.rows):
        for c,cell in enumerate(row.cells):
            cell.width=Inches(963/1440)
            tcw=cell._tc.get_or_add_tcPr().find(qn('w:tcW'));tcw.set(qn('w:w'),'963');tcw.set(qn('w:type'),'dxa')
            n=(c//2)*11+r
            text=('题号' if c%2==0 else '答案') if r==0 else (str(n) if c%2==0 else questions[n-1]['answer'])
            p=cell.paragraphs[0];p.style='AnswerTable';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(text).bold=(r==0 or c%2==1)
            va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');cell._tc.get_or_add_tcPr().append(va)
            if r==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E7E7');cell._tc.get_or_add_tcPr().append(sh)
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t
