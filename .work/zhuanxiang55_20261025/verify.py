"""Check exact models and the finished Word bindings for all 55 questions."""
from collections import Counter
from pathlib import Path
import json,sys
from docx import Document
from PIL import Image
import graphics_a,graphics_b,logic_groups,argument
from content import QUESTIONS,GROUPS
from suite_record import ROOT,DOCX,RECORD,sha

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    reports={name:module.verify() for name,module in [('graphics_a',graphics_a),('graphics_b',graphics_b),('groups',logic_groups),('arguments',argument)]}
    assert [q['n'] for q in QUESTIONS]==list(range(1,56))
    assert Counter(q['category'] for q in QUESTIONS)=={'图形推理':30,'逻辑判断一拖五':15,'逻辑判断加强削弱':10}
    assert Counter(q['kind'] for q in QUESTIONS[45:])=={'加强':5,'削弱':5}
    assert set(GROUPS)=={31,36,41}
    assert all(len(q['options']) in (0,4) and q['answer'] in 'ABCD' and q['analysis'] for q in QUESTIONS)
    assert all(q['image'] for q in QUESTIONS[:30]) and all(not q['image'] for q in QUESTIONS[30:])
    image_hashes=[]
    for q in QUESTIONS[:30]:
        path=ROOT/'assets'/q['image'];assert path.is_file();image_hashes.append(sha(path))
        with Image.open(path) as im:assert im.width==1800 and im.height>=600
    assert len(set(image_hashes))==30,'重复整题图'
    layouts=json.loads((ROOT/'figure_layout.json').read_text(encoding='utf-8'))
    for n in (5,6,7,8):
        layout=layouts[str(n)]
        assert layout['kind']=='two_groups_horizontal'
        assert len(layout['stem_centers'])==6 and len({y for x,y in layout['stem_centers']})==1
    for n in range(22,31):assert layouts[str(n)]['kind']=='classification_3x2'
    doc=Document(DOCX);paras=doc.paragraphs
    qps=[p for p in paras if p.style.name=='Question'];aps=[p for p in paras if p.style.name=='AnswerTitle']
    assert len(qps)==len(aps)==55
    assert [p.text for p in qps]==[q['stem'] for q in QUESTIONS]
    assert [p.text for p in aps]==['答案：'+q['answer'] for q in QUESTIONS]
    text='\n'.join(p.text for p in paras)
    for q in QUESTIONS:
        for s in [q['ask']]+q['options']+q['analysis']:
            if s:assert s in text,(q['n'],s)
    for material in GROUPS.values():
        for s in material:assert s in text
    expected_images=sum(bool(q['image'])+bool(q['explanation_image']) for q in QUESTIONS)
    assert len(doc.inline_shapes)==expected_images
    table=doc.tables[0]
    for n,q in enumerate(QUESTIONS,1):
        r=(n-1)%11+1;c=(n-1)//11*2
        assert table.cell(r,c).text==str(n) and table.cell(r,c+1).text==q['answer']
    assert RECORD.is_file()
    report={'passed':True,'questions':55,'word_sha256':sha(DOCX),'images':expected_images,
            'categories':dict(Counter(q['category'] for q in QUESTIONS)),
            'answers':dict(Counter(q['answer'] for q in QUESTIONS)),
            'graphic_estimated_levels':dict(Counter(q['level'] for q in QUESTIONS[:30])),
            'horizontal_group_questions':[5,6,7,8],
            'model_reports':reports}
    (ROOT/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({key:value for key,value in report.items() if key!='model_reports'},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
