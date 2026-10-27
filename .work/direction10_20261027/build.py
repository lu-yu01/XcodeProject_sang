from pathlib import Path
import sys,json,hashlib,zipfile
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from model import *
sys.path.insert(0,str(ROOT.parent/'advanced_reasoning110_20261026'))
from layout import Numbers,keep,picture

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def build():
 evidence=json.loads((ROOT/'template_reference.json').read_text(encoding='utf-8'))
 assert sha(REF)==evidence['reference_sha256']
 doc=Document(REF);body=doc.element.body
 for e in list(body):
  if e.tag!=qn('w:sectPr'):body.remove(e)
 # Drop only obsolete image relationships; opaque and style parts are preserved.
 for rid,rel in list(doc.part.rels.items()):
  if rel.reltype.endswith('/image'):doc.part.drop_rel(rid)
 doc.core_properties.title='北京风格情境方向图形推理'
 doc.core_properties.subject='十道原创情境方向题及答案图解'
 doc.core_properties.author=''
 numbers=Numbers(doc);nid=numbers.make(left=260,hang=260)
 for page in range(5):
  if page:doc.add_page_break();doc.add_paragraph(f'第{page*2+1}至{page*2+2}题','Heading 1')
  else:
   doc.add_paragraph('北京风格情境方向图形推理','Title')
   doc.add_paragraph('原创练习 十题 附答案与图解','Subtitle')
   doc.add_paragraph('每题选择一个答案。依据画面中的观察方向、景物位置、影向或遮挡关系作答，各题独立。地标和方位标注为解题信息，候选外框不参与判断。')
  for q in Q[page*2:page*2+2]:
   p=doc.add_paragraph(q['stem'],'Stem');numbers.apply(p,nid)
   pic=picture(doc,ROOT/'assets'/f'question-{q["n"]:02}.png',f'第{q["n"]}题的情境示意及候选。',max_width=6.68,max_height=3.5)
   paras=[p,pic]
   if q['n']<=5:
    options=doc.add_paragraph('    '.join(f'{l}. {v}' for l,v in zip('ABCD',q['options'])),'Stem');paras.append(options)
   keep(paras)
   space=doc.add_paragraph();space.paragraph_format.space_after=Pt(4);space.paragraph_format.space_before=Pt(0);space.paragraph_format.line_spacing=1
 for page in range(5):
  doc.add_page_break();doc.add_paragraph('答案与解析' if page==0 else f'第{page*2+1}至{page*2+2}题解析','Heading 1')
  if page==0:doc.add_paragraph('参考答案：'+'    '.join(str(q['n'])+' '+q['answer'] for q in Q))
  for q in Q[page*2:page*2+2]:
   h=doc.add_paragraph(f'第{q["n"]}题 {q["answer"]}','Heading 2')
   p=doc.add_paragraph(q['explanation'],'Explanation');pic=picture(doc,ROOT/'assets'/f'solution-{q["n"]:02}.png',f'第{q["n"]}题方向或顺序解析。',max_width=6.68,max_height=2)
   e=doc.add_paragraph('干扰项辨析：'+q['distractors'],'Explanation');keep([h,p,pic,e])
 assert len(doc.inline_shapes)==20
 assert not list(doc.element.iter(qn('w:pBdr')))
 doc.save(OUT)
 editable={'[Content_Types].xml','docProps/core.xml','word/document.xml','word/numbering.xml','word/_rels/document.xml.rels'}
 with zipfile.ZipFile(REF) as a,zipfile.ZipFile(OUT) as b:
  preserved=[]
  for name in a.namelist():
   if name in editable or name.startswith('word/media/'):continue
   assert a.read(name)==b.read(name),name;preserved.append(name)
 assert Document(OUT).sections[0]._sectPr.xml==Document(REF).sections[0]._sectPr.xml
 (ROOT/'template_fidelity.json').write_text(json.dumps({'reference_sha256':sha(REF),'word_sha256':sha(OUT),'section_identical':True,'preserve_only_identical':preserved},ensure_ascii=False,indent=2),encoding='utf-8')
 (ROOT/'questions.json').write_text(json.dumps(Q,ensure_ascii=False,indent=2),encoding='utf-8')
 (ROOT/'verification.json').write_text(json.dumps(verify(),ensure_ascii=False,indent=2),encoding='utf-8')
 export_public()
 print(json.dumps({'word':str(OUT),'word_sha256':sha(OUT),'images':20,'bytes':OUT.stat().st_size},ensure_ascii=False))
if __name__=='__main__':build()
