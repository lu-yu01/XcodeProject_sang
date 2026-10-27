"""Build one original spatial assembly exercise and its knowledge-base record."""
from pathlib import Path
from collections import Counter
from zipfile import ZipFile
import hashlib,json,sys
from docx import Document
from docx.shared import Mm,Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from layout import style,Numbers,keep,page_field,picture
import draw
from puzzles import QUESTIONS

ROOT=Path(__file__).resolve().parent
PROJECT=ROOT.parent.parent
TITLE='立体拼合专项练习10题'
DOCX=PROJECT/'模式题'/'立体拼合专项练习-10题（含答案解析）.docx'
MD=PROJECT/'知识库'/'套卷索引'/'立体拼合专项练习-10题（含答案解析）.md'

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()

def answer_table(doc):
    t=doc.add_table(rows=2,cols=10);t.autofit=False
    pr=t._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+side)
        for name,val in [('val','single'),('sz','5'),('color','D9D9D9')]:el.set(qn('w:'+name),val)
        b.append(el)
    pr.append(b)
    mar=OxmlElement('w:tblCellMar')
    for side in ('top','bottom','start','end'):
        el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');mar.append(el)
    pr.append(mar)
    for i,q in enumerate(QUESTIONS):
        for row,value in [(0,str(q['n'])),(1,q['answer'])]:
            c=t.cell(row,i);c.width=Inches(.66)
            p=c.paragraphs[0];p.style='Small';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(value).bold=True
            if row==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8E8E8');c._tc.get_or_add_tcPr().append(sh)
    return t

def write_pending_record():
    text='''# 立体拼合专项练习10题 索引与检查记录

## 套卷索引

- Word成品：[立体拼合专项练习10题](../../模式题/立体拼合专项练习-10题（含答案解析）.docx)
- 源稿与构建目录：[assembly10_20261025](../../.work/assembly10_20261025/)。
- 题量：10道原创立体拼合单选题，题目在前，答案与分层拼合解析在后。
- Word版本指纹：`WORDHASH`。

## 生成后检查

- 当前状态：已生成，等待最终数学、显示完整性及逐页视觉检查。
- 难度为命题估计；未进行独立教师盲审或学员试测。
'''
    MD.parent.mkdir(exist_ok=True)
    MD.write_text(text.replace('WORDHASH',sha(DOCX)),encoding='utf-8')

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    draw.draw_all(QUESTIONS)
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(20)
    sec.footer_distance=Mm(9)
    doc.core_properties.title=TITLE;doc.core_properties.subject='立体拼合原创专项练习';doc.core_properties.author='教研资料'
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    style(doc.styles['Normal'],11)
    style(doc.styles['Title'],20,True,0,8,1.1,'黑体')
    style(doc.styles['Heading 1'],14,True,8,6,1.1,'黑体')
    for name,opts in {
        'Question':dict(size=11,before=6,after=5,line=1.25),
        'Figure':dict(size=1,before=3,after=4,line=1),
        'Small':dict(size=9.5,before=0,after=4,line=1.15),
        'AnswerTitle':dict(size=11,bold=True,before=6,after=4,east='黑体'),
        'Point':dict(size=10,bold=True,before=0,after=4),
        'Analysis':dict(size=10.5,before=0,after=3,line=1.15)
    }.items():
        st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);st.base_style=doc.styles['Normal'];style(st,**opts)
    for name in ('AnswerTitle','Point'):doc.styles[name].paragraph_format.keep_with_next=True
    for name in ('Header','Footer'):style(doc.styles[name],8,False,0,0,1)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.add_run('第 ');page_field(footer);footer.add_run(' 页')
    p=doc.add_paragraph(TITLE,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('本练习共10道单选题，练习刚性部件旋转、凹凸互补与多部件立体拼合。','Small')
    intro=[
        '作答约定：各小立方体大小相同；一个部件内部已固定，不可拆开。允许部件整体旋转和平移，不得镜像、变形、重叠或留下空缺。',
        '每个图框内的两幅图是同一物体的两个观察角度，标题注明总块数；灰度仅用于区分面的方向。上方为部件或目标，下方A—D为选项。',
        '姓名：________________　日期：________________　得分：________'
    ]
    for s in intro:doc.add_paragraph(s,'Small')
    nums=Numbers(doc);qid=nums.make();aid=nums.make()
    for i,q in enumerate(QUESTIONS):
        if i:doc.add_page_break()
        block=[]
        p=doc.add_paragraph(q['stem'],'Question');nums.apply(p,qid);block.append(p)
        block.append(picture(doc,draw.ASSETS/f'q{q["n"]:02}.png',f'第{q["n"]}题部件及A至D选项 每个物体含两个视角',6.65,5.45))
        keep(block)
    doc.add_page_break()
    p=doc.add_paragraph('参考答案与解析','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    answer_table(doc)
    doc.add_paragraph('分层图按同一朝向自底向上叠放；甲、乙、丙或选项字样表示各部件占据的小立方体。标为空的格子位于目标之外。','Small')
    for i,q in enumerate(QUESTIONS):
        if i and i%2==0:doc.add_page_break()
        block=[]
        p=doc.add_paragraph('答案：'+q['answer'],'AnswerTitle');nums.apply(p,aid);block.append(p)
        block.append(doc.add_paragraph('考点：'+q['point'],'Point'))
        for j,s in enumerate(q['analysis']):block.append(doc.add_paragraph(('解析：' if j==0 else '')+s,'Analysis'))
        if q.get('except'):
            block.append(doc.add_paragraph('下图展示可拼合选项'+q.get('solution_for','')+'的一个构造；本题所选答案是唯一不能拼合项。','Small'))
        block.append(picture(doc,draw.ASSETS/f'q{q["n"]:02}_solution.png',f'第{q["n"]}题分层拼合构造',6.5,1.65))
        keep(block)
    DOCX.parent.mkdir(exist_ok=True);doc.save(DOCX)
    with ZipFile(DOCX) as z:
        external_images=[e for e in z.namelist() if e.startswith('word/media/')]
    assert len(external_images)==20
    info={'docx':str(DOCX),'word_sha256':sha(DOCX),'question_count':len(QUESTIONS),'embedded_images':len(external_images),
          'answers':dict(Counter(q['answer'] for q in QUESTIONS)),'difficulty_estimates':dict(Counter(q['difficulty'] for q in QUESTIONS))}
    (ROOT/'build_info.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
    write_pending_record()
    print(json.dumps(info,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
