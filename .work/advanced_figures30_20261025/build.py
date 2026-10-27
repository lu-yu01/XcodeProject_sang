"""Build the original advanced 30-item diagram workbook from verified modules."""
from pathlib import Path
from collections import Counter
import json,importlib
from docx import Document
from docx.shared import Mm,Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import drawkit as k
from layout import style,Numbers,keep,page_field,picture

ROOT=Path(__file__).resolve().parent;PROJECT=ROOT.parent.parent
TITLE='国考图形推理进阶专项练习'
DOCX=PROJECT/'模式题'/'国考图形推理进阶专项练习-30题（含答案解析）.docx'
INDEX=PROJECT/'知识库'/'套卷索引'/(DOCX.stem+'.md')
MODULES=['advanced_a','advanced_b','advanced_c']

def data(draw=False):
    questions=[];reports={}
    for name in MODULES:
        mod=importlib.import_module(name)
        if draw:mod.draw_all()
        questions.extend(mod.QUESTIONS);reports[name]=mod.verify()
    questions.sort(key=lambda q:q['n'])
    assert [q['n'] for q in questions]==list(range(1,31))
    for q in questions:
        if isinstance(q['analysis'],str):q['analysis']=[s.strip() for s in q['analysis'].splitlines() if s.strip()]
        assert isinstance(q['analysis'],list) and all(isinstance(s,str) and s.strip() for s in q['analysis'])
        if q.get('ask') in ['从所给的四个选项中，选择最合适的一项。','将六幅图形分成两类，分类正确的是哪一项？']:q['ask']=''
        assert q['answer'] in 'ABCD' and q['analysis'] and q['point']
        assert q['image'] and (k.ASSETS/q['image']).exists(),q['n']
        if q.get('explanation_image'):assert (k.ASSETS/q['explanation_image']).exists(),q['n']
        if q.get('options'):assert len(q['options'])==4,q['n']
    return questions,reports

def table(doc,questions):
    t=doc.add_table(rows=7,cols=10);t.autofit=False
    for col in t.columns:col.width=Inches(.66)
    for r,row in enumerate(t.rows):
        for c,cell in enumerate(row.cells):
            n=(c//2)*6+r
            val=('题号' if c%2==0 else '答案') if r==0 else (str(n) if c%2==0 else questions[n-1]['answer'])
            p=cell.paragraphs[0];p.style='AnswerTable';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(val).bold=(r==0 or c%2==1)
            if r==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E7E7');cell._tc.get_or_add_tcPr().append(sh)
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
    return t

def main():
    questions,reports=data(draw=True)
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(20)
    sec.header_distance=Mm(8);sec.footer_distance=Mm(9)
    doc.core_properties.title=TITLE;doc.core_properties.subject='原创30题及答案解析';doc.core_properties.author='教研资料'
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    style(doc.styles['Normal']);style(doc.styles['Title'],20,True,0,7,1.1,'黑体')
    style(doc.styles['Subtitle'],10.5,False,0,5,1.15)
    style(doc.styles['Heading 1'],15,True,10,7,1.1,'黑体');doc.styles['Heading 1'].paragraph_format.keep_with_next=True
    style(doc.styles['Heading 2'],12,True,7,5,1.1,'黑体');doc.styles['Heading 2'].paragraph_format.keep_with_next=True
    params={'Question':dict(size=11,before=8,after=4),
            'Ask':dict(size=10.5,before=0,after=3),
            'Option':dict(size=11,before=0,after=2),
            'Figure':dict(size=1,before=2,after=4,line=1),
            'AnswerTitle':dict(size=11,bold=True,before=7,after=3,east='黑体'),
            'Point':dict(size=10,bold=True,before=0,after=2),
            'Analysis':dict(size=10.5,before=0,after=3,line=1.15),
            'Small':dict(size=9.5,before=0,after=5,line=1.15),
            'AnswerTable':dict(size=10,before=0,after=3,line=1.1)}
    for name,params_ in params.items():
        st=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);st.base_style=doc.styles['Normal'];style(st,**params_)
    for name in ('AnswerTitle','Point'):doc.styles[name].paragraph_format.keep_with_next=True
    for name in ('Header','Footer'):style(doc.styles[name],8,False,0,0,1)
    sec.header.paragraphs[0].text=TITLE
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('第 ');page_field(p);p.add_run(' 页')
    for text,s in [(TITLE,'Title'),('原创题目册 · 30题','Subtitle')]:
        p=doc.add_paragraph(text,s);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('共30道单项选择题，每题只有一个正确答案。建议用时45分钟。','Small')
    doc.add_paragraph('姓名：________________　日期：________________　得分：________','Small')
    doc.add_paragraph('各图外统一增设的细线框仅用于排版，不参与规律判断；图形自身的边线仍属于图形。','Small')
    nums=Numbers(doc);qid=nums.make();aid=nums.make()
    doc.add_paragraph('图形推理','Heading 1')
    for q in questions:
        block=[];p=doc.add_paragraph(q['stem'],'Question');nums.apply(p,qid);block.append(p)
        if q.get('ask'):block.append(doc.add_paragraph(q['ask'],'Ask'))
        block.append(picture(doc,k.ASSETS/q['image'],f"第{q['n']}题题图与选项",6.6,4.25))
        if q.get('options'):
            oid=nums.make('upperLetter',left=450,hang=300)
            for opt in q['options']:
                p=doc.add_paragraph(opt,'Option');nums.apply(p,oid);block.append(p)
        keep(block)
    p=doc.add_paragraph('参考答案与解析','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.page_break_before=True
    doc.add_paragraph('答案速查','Heading 2');table(doc,questions)
    for q in questions:
        block=[];p=doc.add_paragraph('答案：'+q['answer'],'AnswerTitle');nums.apply(p,aid);block.append(p)
        block.append(doc.add_paragraph('考点：'+q['point'],'Point'))
        for i,s in enumerate(q['analysis']):block.append(doc.add_paragraph(('解析：' if i==0 else '')+s,'Analysis'))
        if q.get('explanation_image'):block.append(picture(doc,k.ASSETS/q['explanation_image'],f"第{q['n']}题解析辅助图",6.25,3.0))
        keep(block)
    DOCX.parent.mkdir(exist_ok=True);doc.save(DOCX)
    (ROOT/'questions.json').write_text(json.dumps(questions,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'model_verification.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'figure_layout.json').write_text(json.dumps(k.LAYOUTS,ensure_ascii=False,indent=2),encoding='utf-8')
    meta={'word':str(DOCX),'questions':30,'answers':dict(Counter(q['answer'] for q in questions)),
          'inline_images':len(doc.inline_shapes),'index':str(INDEX)}
    (ROOT/'build_info.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    INDEX.write_text(f'# {DOCX.stem} 索引与检查记录\n\n- Word：[成品](../../模式题/{DOCX.name})。\n- 工作目录：[本卷源稿](../../.work/advanced_figures30_20261025/)。\n- 30题，题目在前，答案与解析在后。\n- 状态：模型与成品正在复核，尚未最终验收。\n',encoding='utf-8')
    print(json.dumps(meta,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
