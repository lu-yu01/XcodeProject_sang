from pathlib import Path
import sys,json,hashlib,copy
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Pt,Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
ROOT=Path(__file__).resolve().parent;PROJECT=ROOT.parents[1]
sys.path.insert(0,str(ROOT.parent/'advanced_reasoning110_20261026'))
from layout import Numbers,keep,picture
from model import SYMBOLS,TEXT
REF=PROJECT/'模式题'/'北京风格图形推理一拖三-1组3题（含答案解析）.docx'
OUT=PROJECT/'模式题'/'北京风格图形操作进阶练习-20组60题（含答案解析）.docx'
ASSETS=ROOT/'assets';ASSETS.mkdir(exist_ok=True)
GROUPS=json.loads((ROOT/'questions.json').read_text(encoding='utf-8'))
def font(s,bold=False):return ImageFont.truetype('C:/Windows/Fonts/simhei.ttf' if bold else 'C:/Windows/Fonts/simsun.ttc',s)
def center(d,xy,t,f):d.text(xy,t,font=f,fill='black',anchor='mm')
def grid(d,a,x,y,s=240):
 u=s/4
 for r,row in enumerate(a):
  for c,v in enumerate(row):
   x0,y0=x+c*u,y+r*u;cx,cy=x0+u/2,y0+u/2
   ps=[[(x0,y0),(x0+u,y0),(cx,cy)],[(x0+u,y0),(x0+u,y0+u),(cx,cy)],[(x0+u,y0+u),(x0,y0+u),(cx,cy)],[(x0,y0+u),(x0,y0),(cx,cy)]]
   for i,p in enumerate(ps):
    if v>>i&1:d.polygon(p,fill='black')
 for i in range(5):
  d.line((x+i*u,y,x+i*u,y+s),fill='black',width=3);d.line((x,y+i*u,x+s,y+i*u),fill='black',width=3)
def glyph(d,n,x,y,s=30):
 symbol=SYMBOLS[n];shape=symbol[0];x-=12
 filled=shape in '●◆▲■';fill='black' if filled else 'white'
 if shape in '●○':d.ellipse((x-s/2,y-s/2,x+s/2,y+s/2),fill=fill,outline='black',width=3)
 elif shape in '◆◇':d.polygon([(x,y-s*.67),(x+s*.58,y),(x,y+s*.67),(x-s*.58,y)],fill=fill,outline='black',width=3)
 elif shape in '▲△':d.polygon([(x,y-s*.61),(x+s*.65,y+s*.5),(x-s*.65,y+s*.5)],fill=fill,outline='black',width=3)
 else:d.rectangle((x-s/2,y-s/2,x+s/2,y+s/2),fill=fill,outline='black',width=3)
 center(d,(x+34,y+2),symbol[1],font(36,True))
def arrow(d,x1,y,x2):
 d.line((x1,y,x2,y),fill='black',width=3);d.polygon([(x2,y),(x2-13,y-8),(x2-13,y+8)],fill='black')
def draw_question(q):
 im=Image.new('RGB',(1800,365),'white');d=ImageDraw.Draw(im);grid(d,q['source'],25,30)
 arrow(d,280,150,304)
 for j,n in enumerate(q['seq']):
  x=340+j*115;glyph(d,n,x,150)
  arrow(d,x+43,150,x+77 if j<2 else 657)
 left,top,col,h=700,14,265,272
 d.rectangle((left,top,left+4*col,top+h),outline='black',width=3)
 for j,a in enumerate(q['options']):
  if j:d.line((left+j*col,top,left+j*col,top+h),fill='black',width=3)
  grid(d,a,left+j*col+13,30);center(d,(left+j*col+col/2,322),'ABCD'[j],font(43,True))
 im.save(ASSETS/f'question-{q["n"]:02}.png',dpi=(300,300))
def draw_solution(q):
 im=Image.new('RGB',(1800,330),'white');d=ImageDraw.Draw(im);xs=[95,550,1005,1460]
 for j,(a,x) in enumerate(zip(q['states'],xs)):
  grid(d,a,x,18);center(d,(x+120,296),'原图' if j==0 else f'第{j}步',font(34))
  if j<3: glyph(d,q['seq'][j],(x+240+xs[j+1])/2,100);arrow(d,x+275,175,xs[j+1]-40)
 im.save(ASSETS/f'solution-{q["n"]:02}.png',dpi=(300,300))
def rule_table(doc,g):
 t=doc.add_table(rows=6,cols=4);t.autofit=False;widths=[.50,2.84,.50,2.84]
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 pr=t._tbl.tblPr;tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'9619')
 borders=OxmlElement('w:tblBorders')
 for side in ('top','left','bottom','right','insideH','insideV'):
  e=OxmlElement('w:'+side)
  for k,v in [('val','single'),('sz','5'),('color','000000')]:e.set(qn('w:'+k),v)
  borders.append(e)
 pr.append(borders)
 mar=OxmlElement('w:tblCellMar')
 for side,val in [('top',65),('bottom',65),('start',60),('end',60)]:
  e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');mar.append(e)
 pr.append(mar)
 # Keep similar operators close rather than introduce false lookalike glyphs.
 for r,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for h in range(2):
   n=g['rules'][r*2+h]
   for j,text in enumerate((SYMBOLS[n],TEXT[n])):
    c=row.cells[h*2+j];c.width=Inches(widths[h*2+j]);p=c.paragraphs[0];p.style='Rules'
    if j==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run(text)
    if j==0:run.font.size=Pt(13)
    va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');c._tc.get_or_add_tcPr().append(va)
 return t
def build():
 baseline=ROOT/'baseline_docx.json'
 if not baseline.exists():baseline.write_text(json.dumps({str(p.relative_to(PROJECT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in (PROJECT/'模式题').glob('*.docx')},ensure_ascii=False,indent=2),encoding='utf-8')
 for g in GROUPS:
  for q in g['questions']:draw_question(q);draw_solution(q)
 # The previous approved two-page document remains the formatting authority.
 doc=Document(REF);body=doc._element.body
 for child in list(body):
  if child.tag!=qn('w:sectPr'):body.remove(child)
 nums=Numbers(doc);nid=nums.make(left=260,hang=260)
 for gi,g in enumerate(GROUPS):
  if gi:doc.add_page_break()
  if gi==0:
   doc.add_paragraph('北京风格图形操作进阶练习','Title')
   doc.add_paragraph('20组60题  每组12条规则','Subtitle')
  doc.add_paragraph(f'第{g["group"]}组  第{gi*3+1}至{gi*3+3}题','Heading 1')
  doc.add_paragraph('本组三题共用下表，各题独立。从左到右依次操作，每步以上一步结果为对象；上下左右以当前图为准。')
  doc.add_paragraph('符号右侧数字也是标识的一部分，含义全卷一致。四象限为四个2×2区域；“四角小格”指整图最外侧四个角格。')
  rule_table(doc,g)
  p=doc.add_paragraph('镜像改变内部朝向；直接换位、平移复制保持内部朝向。反色只改变填充，格线不变。');p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(5)
  for q in g['questions']:
   p=doc.add_paragraph('对左侧图形依次操作后，可以得到右侧哪个图形？','Stem');nums.apply(p,nid)
   p2=picture(doc,ASSETS/f'question-{q["n"]:02}.png',f'第{q["n"]}题输入图、三步操作和四候选。',max_width=6.68,max_height=2);keep([p,p2])
 for gi,g in enumerate(GROUPS):
  doc.add_page_break()
  if gi==0:doc.add_paragraph('答案与解析','Heading 1')
  doc.add_paragraph(f'第{g["group"]}组  第{gi*3+1}至{gi*3+3}题','Heading 1')
  doc.add_paragraph('参考答案：'+'    '.join(str(q['n'])+' '+q['answer'] for q in g['questions']))
  for q in g['questions']:
   h=doc.add_paragraph(f'第{q["n"]}题','Heading 2')
   text='；'.join(f'第{i+1}步（{SYMBOLS[n]}）：{TEXT[n]}' for i,n in enumerate(q['seq']))+f'。最终结果对应{q["answer"]}项。'
   p=doc.add_paragraph(text,'Explanation');keep([h,p])
   picture(doc,ASSETS/f'solution-{q["n"]:02}.png',f'第{q["n"]}题的原图和三步中间结果。',max_width=6.68,max_height=2)
   wrong=[f'{e["letter"]}项：{e["reason"]}。' for e in q['option_explanations'] if not e['correct']]
   doc.add_paragraph('干扰项辨析：'+''.join(wrong),'Explanation')
 doc.core_properties.title='北京风格图形操作进阶练习'
 doc.core_properties.subject='20组60题 每组12项规则及分步解析'
 doc.save(OUT)
 print(json.dumps({'word':str(OUT),'bytes':OUT.stat().st_size,'assets':len(list(ASSETS.glob('*.png')))},ensure_ascii=False))
if __name__=='__main__':build()
