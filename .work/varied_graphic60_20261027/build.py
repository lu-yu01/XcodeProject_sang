import sys,copy,json,zipfile
from docx import Document
from docx.shared import Pt,Twips
from docx.oxml.ns import qn
from docx.table import Table
from model import *
sys.path.insert(0,str(ROOT.parent/'advanced_reasoning110_20261026'))
from layout import Numbers,keep,picture

def rule_table(doc,template,rules):
 element=copy.deepcopy(template);doc.element.body.insert(len(doc.element.body)-1,element)
 table=Table(element,doc._body)
 widths=[840,3969,840,3970]
 for col,width in zip(table.columns,widths):col.width=Twips(width)
 while len(table.rows)>(len(rules)+1)//2:table._tbl.remove(table.rows[-1]._tr)
 for i,row in enumerate(table.rows):
  for cell,width in zip(row.cells,widths):cell.width=Twips(width)
  for side in range(2):
   ix=2*i+side;codecell,textcell=row.cells[2*side:2*side+2]
   if ix>=len(rules):codecell.text='';textcell.text='';continue
   r=rules[ix];codecell.text=r['code'];textcell.text=r['text']
   for p in codecell.paragraphs:
    p.style='Rules';p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
    for run in p.runs:
     run.font.name='Segoe UI Symbol';run.font.size=Pt(12)
     rf=run._r.get_or_add_rPr().get_or_add_rFonts()
     for key in ['ascii','hAnsi','eastAsia']:rf.set(qn('w:'+key),'Segoe UI Symbol')
   for p in textcell.paragraphs:
    p.style='Rules';p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
 return table

def build():
 evidence=json.loads((ROOT/'template_reference.json').read_text(encoding='utf-8'));assert sha(REF)==evidence['reference_sha256']
 groups=json.loads((ROOT/'groups.json').read_text(encoding='utf-8'));questions=json.loads((ROOT/'questions.json').read_text(encoding='utf-8'))
 doc=Document(REF);template=copy.deepcopy(doc.tables[0]._tbl);body=doc.element.body
 for e in list(body):
  if e.tag!=qn('w:sectPr'):body.remove(e)
 for rid,rel in list(doc.part.rels.items()):
  if rel.reltype.endswith('/image'):doc.part.drop_rel(rid)
 doc.core_properties.title='北京风格图形操作多样式专项练习';doc.core_properties.subject='二十组原创一拖三图形操作题及答案图解';doc.core_properties.author=''
 numbers=Numbers(doc);nid=numbers.make(left=260,hang=260)
 for i,g in enumerate(groups):
  if i:doc.add_page_break()
  else:
   doc.add_paragraph('北京风格图形操作多样式专项练习','Title')
   doc.add_paragraph('原创练习 二十组六十题 附答案与分步图解','Subtitle')
   doc.add_paragraph('每组共用本组操作表，三题独立起算。按操作链从左到右依次执行，每一步以上一步结果为对象。相同符号在不同组只按本组表解释；候选外框及A至D不参加操作。','Rules')
  doc.add_paragraph(f'第{g["id"]}组 {g["title"].replace("、", " ")}','Heading 1')
  doc.add_paragraph(g['intro'],'Rules')
  rule_table(doc,template,g['rules'])
  # Separate the table from the first stem without introducing a large blank.
  p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(1);p.add_run().font.size=Pt(1)
  for q in (q for q in questions if q['group']==g['id']):
   p=doc.add_paragraph('对左侧图形依次操作后，可以得到右侧哪个图形？','Stem');numbers.apply(p,nid)
   pic=picture(doc,ROOT/'assets'/f'question-{q["n"]:02}.png',f'第{q["n"]}题 输入图 操作链及四候选',max_width=6.68,max_height=1.55);keep([p,pic])
 for i,g in enumerate(groups):
  doc.add_page_break();doc.add_paragraph('答案与解析' if i==0 else f'第{g["id"]}组解析','Heading 1')
  if i==0:doc.add_paragraph('各组解析按原题顺序排列。分步图自左向右展示当前状态；操作代码仍以相应题组的公开表为准。','Rules')
  qs=[q for q in questions if q['group']==g['id']]
  doc.add_paragraph(f'第{g["id"]}组 {g["title"].replace("、", " ")}','Heading 2')
  doc.add_paragraph('参考答案：'+'    '.join(f'{q["n"]} {q["answer"]}' for q in qs),'Explanation')
  rules={r['id']:r for r in g['rules']}
  for q in qs:
   h=doc.add_paragraph(f'第{q["n"]}题 {q["answer"]}','Heading 2')
   text='；'.join(f'{j+1} {rules[r]["code"]}：{rules[r]["text"].rstrip("。；")}' for j,r in enumerate(q['chain']))+'。'
   p=doc.add_paragraph(text,'Explanation')
   pic=picture(doc,ROOT/'assets'/f'solution-{q["n"]:02}.png',f'第{q["n"]}题的原图及逐步结果',max_width=6.68,max_height=1.5)
   text='  '.join(f'{letter}项：{o["explanation"]}' for letter,o in zip('ABCD',q['options']) if letter!=q['answer'])
   e=doc.add_paragraph('干扰项辨析：'+text,'Explanation');keep([h,p,pic,e])
 assert len(doc.inline_shapes)==120 and len(doc.tables)==20
 assert not list(doc.element.iter(qn('w:pBdr')))
 doc.save(OUT)
 editable={'[Content_Types].xml','docProps/core.xml','word/document.xml','word/numbering.xml','word/_rels/document.xml.rels'}
 with zipfile.ZipFile(REF) as a,zipfile.ZipFile(OUT) as b:
  preserved=[]
  for name in a.namelist():
   if name in editable or name.startswith('word/media/'):continue
   assert a.read(name)==b.read(name),name;preserved.append(name)
 assert Document(OUT).sections[0]._sectPr.xml==Document(REF).sections[0]._sectPr.xml
 dump(ROOT/'template_fidelity.json',{'reference_sha256':sha(REF),'word_sha256':sha(OUT),'section_identical':True,'preserve_only_identical':preserved})
 print(json.dumps({'word':str(OUT),'word_sha256':sha(OUT),'pictures':120,'bytes':OUT.stat().st_size},ensure_ascii=False))
if __name__=='__main__':build()
