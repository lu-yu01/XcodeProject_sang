from pathlib import Path
import json,hashlib,re,zipfile,datetime
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
from model import QUESTIONS,SYMBOLS,RULES,verify
ROOT=Path(__file__).resolve().parent;PROJECT=ROOT.parents[1]
OUT=PROJECT/'模式题'/'北京风格图形推理一拖三-1组3题（含答案解析）.docx'
INDEX=PROJECT/'知识库'/'套卷索引'/(OUT.stem+'.md')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read_grid(im,a,x,y,s=240):
 # Four interior triangle centroids, independently read from the rendered pixels.
 offsets=[(.5,1/6),(5/6,.5),(.5,5/6),(1/6,.5)]
 for r,row in enumerate(a):
  for c,v in enumerate(row):
   for bit,(dx,dy) in enumerate(offsets):
    rgb=im.getpixel((round(x+(c+dx)*s/4),round(y+(r+dy)*s/4)))
    actual=max(rgb)<30
    assert actual==bool(v>>bit&1),(r,c,bit,rgb,v)
 samples=4*4*4
 return samples
def check():
 model=verify();sample_count=0
 for q in QUESTIONS:
  with Image.open(ROOT/'assets'/f'question-{q["n"]:02}.png') as im:
   sample_count+=read_grid(im,q['source'],25,30)
   for j,a in enumerate(q['options']):sample_count+=read_grid(im,a,700+j*265+13,30)
  with Image.open(ROOT/'assets'/f'solution-{q["n"]:02}.png') as im:
   for j,a in enumerate(q['states']):sample_count+=read_grid(im,a,[95,550,1005,1460][j],18)
 assets={sha(p) for p in (ROOT/'assets').glob('*.png')}
 d=Document(OUT);inline_checks=[]
 assert len(d.inline_shapes)==6
 for shape in d.inline_shapes:
  rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
  blob=d.part.related_parts[rid].blob
  assert hashlib.sha256(blob).hexdigest() in assets
  import io
  with Image.open(io.BytesIO(blob)) as im:ratio=im.width/im.height
  err=abs(float(shape.width)/float(shape.height)-ratio)
  assert err<.00002
  inline_checks.append({'image_hash':hashlib.sha256(blob).hexdigest(),'ratio_error':err})
 assert len(d.tables)==1 and len(d.tables[0].rows)==4
 celltext=''.join(c.text for row in d.tables[0].rows for c in row.cells)
 for name,text in RULES:assert SYMBOLS[name] in celltext and text in celltext
 assert not list(d.element.iter(qn('w:pBdr')))
 assert not list(d.styles.element.iter(qn('w:pBdr')))
 texts=[p.extract_text() for p in PdfReader(ROOT/'preview.pdf').pages]
 assert len(texts)==2
 assert '答案与解析' not in texts[0] and '答案与解析' in texts[1]
 for i in range(1,4):assert f'{i}. 对左侧图形' in texts[0]
 for s in SYMBOLS.values():assert s in texts[0]
 old=json.loads((ROOT/'baseline_docx.json').read_text(encoding='utf-8'))
 assert all((PROJECT/p).exists() and sha(PROJECT/p)==h for p,h in old.items())
 blind=json.loads((ROOT/'blind_review.json').read_text(encoding='utf-8'))
 assert blind['answers']=={str(q['n']):q['answer'] for q in QUESTIONS}
 visual=json.loads((ROOT/'visual_review.json').read_text(encoding='utf-8'))
 assert visual['word_sha256']==sha(OUT) and visual['pdf_sha256']==sha(ROOT/'preview.pdf')
 assert len(visual['pages'])==2 and all(x['actually_opened'] and x['passed'] for x in visual['pages'])
 assert INDEX.exists()
 links=[]
 for file in [INDEX,PROJECT/'知识库'/'北京一拖三图形操作题命题规则.md']:
  for link in re.findall(r'\]\(([^)]+)\)',file.read_text(encoding='utf-8')):
   if '://' not in link:
    target=(file.parent/link).resolve();assert target.exists(),str(target);links.append(str(target))
 report={'passed':True,'created_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),
  'word':str(OUT),'word_sha256':sha(OUT),'pdf_sha256':sha(ROOT/'preview.pdf'),'pdf_pages':2,
  'questions':3,'shared_rule_count':8,'answers':{q['n']:q['answer'] for q in QUESTIONS},
  'model_verification':model,'png_interior_samples':sample_count,'embedded_figures':inline_checks,
  'existing_docx_unchanged':len(old),'knowledge_links_checked':len(links),
  'blind_review':'3/3与独立试解一致，详见blind_review.json',
  'visual_review':'最终2页逐页实际打开，详见visual_review.json',
  'limits':['难度为命题估计','未进行教师独立审题或学员试测','本地2026章年份未获独立核准，成品采用北京风格命名']}
 (ROOT/'final_check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps({k:report[k] for k in ['passed','pdf_pages','questions','word_sha256','png_interior_samples','existing_docx_unchanged']},ensure_ascii=False))
if __name__=='__main__':check()
