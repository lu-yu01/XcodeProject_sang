from pathlib import Path
import sys,json,hashlib,math
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
PROJECT=ROOT.parents[1]
sys.path.insert(0,str(ROOT.parent/'advanced_reasoning110_20261026'))
from layout import style,Numbers,keep,picture
from model import QUESTIONS,SYMBOLS,RULES,verify
ASSETS=ROOT/'assets';ASSETS.mkdir(exist_ok=True)
OUT=PROJECT/'模式题'/'北京风格图形推理一拖三-1组3题（含答案解析）.docx'
BASE=ROOT/'baseline_docx.json'
if not BASE.exists():
 BASE.write_text(json.dumps({str(p.relative_to(PROJECT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in (PROJECT/'模式题').glob('*.docx')},ensure_ascii=False,indent=2),encoding='utf-8')

def font(size,bold=False):
 return ImageFont.truetype('C:/Windows/Fonts/simhei.ttf' if bold else 'C:/Windows/Fonts/simsun.ttc',size)
def center(d,xy,t,f,fill='black'):
 d.text(xy,t,font=f,fill=fill,anchor='mm')
def grid(d,a,x,y,s=240):
 u=s/4
 for r,row in enumerate(a):
  for c,v in enumerate(row):
   x0,y0=x+c*u,y+r*u;cx,cy=x0+u/2,y0+u/2
   pts=[[(x0,y0),(x0+u,y0),(cx,cy)],[(x0+u,y0),(x0+u,y0+u),(cx,cy)],[(x0+u,y0+u),(x0,y0+u),(cx,cy)],[(x0,y0+u),(x0,y0),(cx,cy)]]
   for i,p in enumerate(pts):
    if v>>i&1:d.polygon(p,fill='black')
 for i in range(5):
  d.line((x+i*u,y,x+i*u,y+s),fill='black',width=3)
  d.line((x,y+i*u,x+s,y+i*u),fill='black',width=3)
def glyph(d,name,x,y,s=36):
 filled=name in ('CW','LR','IT','CL');fill='black' if filled else 'white'
 # Icons are geometric, so rule glyphs and operation chains stay legible.
 if name in ('CW','CCW'):d.ellipse((x-s/2,y-s/2,x+s/2,y+s/2),fill=fill,outline='black',width=3)
 elif name in ('LR','UD'):d.polygon([(x,y-s*.67),(x+s*.58,y),(x,y+s*.67),(x-s*.58,y)],fill=fill,outline='black',width=3)
 elif name in ('IT','IR'):d.polygon([(x,y-s*.61),(x+s*.65,y+s*.5),(x-s*.65,y+s*.5)],fill=fill,outline='black',width=3)
 else:d.rectangle((x-s/2,y-s/2,x+s/2,y+s/2),fill=fill,outline='black',width=3)
def arrow(d,x1,y,x2):
 d.line((x1,y,x2,y),fill='black',width=3)
 d.polygon([(x2,y),(x2-13,y-8),(x2-13,y+8)],fill='black')
def question(q):
 im=Image.new('RGB',(1800,365),'white');d=ImageDraw.Draw(im)
 grid(d,q['source'],25,30)
 for j,name in enumerate(q['seq']):
  x=350+j*105;glyph(d,name,x,150,39)
  arrow(d,x+31,150,x+70)
 arrow(d,280,150,308)
 # A single enclosing option frame, with labels outside underneath.
 left,top=700,14;col=265;height=272
 d.rectangle((left,top,left+col*4,top+height),outline='black',width=3)
 for j,a in enumerate(q['options']):
  if j:d.line((left+j*col,top,left+j*col,top+height),fill='black',width=3)
  grid(d,a,left+j*col+13,top+16)
  center(d,(left+j*col+col/2,322),'ABCD'[j],font(43,True))
 path=ASSETS/f'question-{q["n"]:02}.png';im.save(path,dpi=(300,300));return path
def solution(q):
 im=Image.new('RGB',(1800,330),'white');d=ImageDraw.Draw(im)
 xs=[95,550,1005,1460]
 for j,(a,x) in enumerate(zip(q['states'],xs)):
  grid(d,a,x,18)
  center(d,(x+120,296),'原图' if j==0 else f'第{j}步',font(34))
  if j<3:
   m=(x+240+xs[j+1])/2;glyph(d,q['seq'][j],m,100,38)
   arrow(d,x+275,175,xs[j+1]-40)
 path=ASSETS/f'solution-{q["n"]:02}.png';im.save(path,dpi=(300,300));return path

def table_rules(doc):
 t=doc.add_table(rows=4,cols=4);t.autofit=False
 widths=[.34,3.0,.34,3.0]
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 pr=t._tbl.tblPr
 tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'9619')
 borders=OxmlElement('w:tblBorders')
 for side in ('top','left','bottom','right','insideH','insideV'):
  e=OxmlElement('w:'+side)
  for k,v in [('val','single'),('sz','5'),('color','000000')]:e.set(qn('w:'+k),v)
  borders.append(e)
 pr.append(borders)
 mar=OxmlElement('w:tblCellMar')
 for side,value in [('top',75),('bottom',75),('start',65),('end',65)]:
  e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');mar.append(e)
 pr.append(mar)
 for r,row in enumerate(t.rows):
  cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
  for h in range(2):
   name,txt=RULES[r*2+h]
   for j,content in enumerate((SYMBOLS[name],txt)):
    c=row.cells[h*2+j];c.width=Inches(widths[h*2+j]);p=c.paragraphs[0];p.style='Rules'
    if j==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run(content)
    if j==0:run.font.size=Pt(14)
    va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');c._tc.get_or_add_tcPr().append(va)
 return t

def build():
 for q in QUESTIONS:question(q);solution(q)
 doc=Document();sec=doc.sections[0]
 sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.65);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
 for st in list(doc.styles):
  if st.type==1:
   for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
 style(doc.styles['Normal'],10.5,after=4,line=1.15)
 style(doc.styles['Title'],19,True,after=5,line=1.1,east='黑体')
 style(doc.styles['Subtitle'],10.5,after=8,line=1.1)
 style(doc.styles['Heading 1'],13,True,before=4,after=6,line=1.1,east='黑体')
 style(doc.styles['Heading 2'],11.5,True,before=7,after=4,line=1.1,east='黑体')
 for name,size,after,line in [('Figure',1,4,1.0),('Rules',10,0,1.08),('Stem',10.5,3,1.08),('Explanation',10.5,4,1.15)]:
  if name not in doc.styles:doc.styles.add_style(name,1)
  style(doc.styles[name],size,after=after,line=line)
 doc.core_properties.title='北京风格图形推理一拖三'
 doc.core_properties.subject='原创一组3题及分步图解'
 doc.core_properties.author=''
 p=doc.add_paragraph('北京风格图形推理一拖三','Title')
 doc.add_paragraph('原创样题  一组3题','Subtitle')
 doc.add_paragraph('三题共用下表，但各题独立。对每题左侧图形，按符号从左到右依次操作，选择最终得到的图形。')
 doc.add_paragraph('每一步都以上一步的结果为对象；“上、下、左、右”均以当时图形为准。反色只改变黑白填充，方格线不变。')
 doc.add_paragraph('共用操作规则','Heading 1');table_rules(doc)
 p=doc.add_paragraph('注：镜像为翻折后的效果；局部镜像只在指定半区内部进行。');p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(5)
 nums=Numbers(doc);nid=nums.make(left=260,hang=260)
 for q in QUESTIONS:
  p=doc.add_paragraph('对左侧图形依次操作后，可以得到右侧哪个图形？','Stem');nums.apply(p,nid)
  pic=picture(doc,ASSETS/f'question-{q["n"]:02}.png',f'第{q["n"]}题：输入图、三个操作符号与四个选项。',max_width=6.68,max_height=2)
  keep([p,pic])
 doc.add_page_break()
 doc.add_paragraph('答案与解析','Heading 1')
 doc.add_paragraph('参考答案：1 B    2 D    3 A')
 summaries=[
  '先整体左右镜像，再整体顺时针旋转90°，最后只将当前图形的上半部分反色。旋转完成后要重新定位上下，结果对应B项。',
  '先保留左半部分，用它的左右镜像重建右半；再将得到的完整图形逆时针旋转90°；最后只对当前图形的右半部分反色，结果对应D项。',
  '先只在上半部分内部做左右镜像，下半部分不变；再将全图上下镜像；最后将全图顺时针旋转90°，结果对应A项。']
 for q,summary in zip(QUESTIONS,summaries):
  h=doc.add_paragraph(f'第{q["n"]}题','Heading 2')
  p=doc.add_paragraph(summary,'Explanation');keep([h,p])
  picture(doc,ASSETS/f'solution-{q["n"]:02}.png',f'第{q["n"]}题：原图及三步变换结果。',max_width=6.68,max_height=2)
  wrong=[f'{e["letter"]}项：{e["reason"]}。' for e in q['option_explanations'] if not e['correct']]
  doc.add_paragraph('干扰项辨析：'+''.join(wrong),'Explanation')
  doc.add_paragraph('考查重点：'+q['point']+'。','Explanation')
 # Cover default template borders as well as direct formatting.
 for b in list(doc.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
 doc.save(OUT)
 (ROOT/'verification.json').write_text(json.dumps(verify(),ensure_ascii=False,indent=2),encoding='utf-8')
 (ROOT/'questions.json').write_text(json.dumps(QUESTIONS,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps({'word':str(OUT),'source_images':len(list(ASSETS.glob('*.png'))),'bytes':OUT.stat().st_size},ensure_ascii=False))
if __name__=='__main__':build()
