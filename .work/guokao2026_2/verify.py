"""Recheck paper structure, Word content and computable figure answers."""
from __future__ import annotations

import json
from collections import Counter

from docx import Document

import figures as f
from content import COMMON, GRAPHIC, QUESTIONS
from logic_draft import verify_common
from suite_record import DOCX, RECORD, ROOT, sha256


def assert_circular_state(candidate, expected):
    return candidate[0] == expected[0] and candidate[1] == expected[1]


def q8_rotation_orbit(state):
    orbit = set()
    for _ in range(4):
        orbit.add(f.q8_canonical(state))
        state = f.q8_transform(state, "rotate")
    return orbit


def verify_figures():
    assert sum(assert_circular_state(x, f.chord_state(5)) for x in f.Q1_OPTIONS) == 1
    assert f.Q1_OPTIONS[2] == f.chord_state(5)
    assert f.Q2_OPTIONS[1] == f.Q2_CORRECT
    assert len({(frozenset(c), p) for c, p in f.Q2_OPTIONS}) == 4
    assert f.Q3_SIGNATURES == [(2, 4)] * 3 + [(3, 2)] * 2
    assert f.Q3_OPTION_SIGNATURES == [(3, 4), (2, 2), (3, 0), (3, 2)]
    assert f.Q4_OPTIONS[0] == f.operation(*f.Q4_ROWS[2])
    assert len({frozenset(x) for x in f.Q4_OPTIONS}) == 4
    assert all(0 <= y <= 1 for _, y, _ in f.Q5_SECTION)
    assert [f.q6_works(x) for x in f.Q6_OPTIONS] == [False, True, False, False]
    assert [w is not None for w in f.Q7_WITNESSES] == [False, False, False, True]
    assert sum(f.Q7_WITNESSES[3].values()) == f.Q7_TOTAL == 12
    first, second = q8_rotation_orbit(f.Q8_STATES[0]), q8_rotation_orbit(f.Q8_STATES[1])
    assert len(first) == len(second) == 4 and not first & second
    assert [f.q8_canonical(x) in first for x in f.Q8_STATES] == [True, False] * 3
    assert f.Q9_PERIMETERS == [12, 14, 12, 14, 12, 14]
    assert f.Q10_VALID == [True, False, True, True, False, False]
    assert [q["answer"] for q in GRAPHIC] == list("CBDACBDACB")
    assert all((f.ASSETS / q["image"]).is_file() for q in GRAPHIC)


def verify_docx():
    assert DOCX.is_file() and RECORD.is_file()
    doc = Document(DOCX)
    para = doc.paragraphs
    question_paras = [p for p in para if p.style.name == "Question"]
    answer_paras = [p for p in para if p.style.name == "AnswerTitle"]
    assert len(question_paras) == len(answer_paras) == len(QUESTIONS) == 35
    assert [p.text for p in question_paras] == [q["stem"] for q in QUESTIONS]
    assert [p.text for p in answer_paras] == ["答案：" + q["answer"] for q in QUESTIONS]
    assert len(doc.inline_shapes) == 10
    all_text = "\n".join(p.text for p in para)
    for line in COMMON:
        assert line in all_text, line
    for q in QUESTIONS:
        if q["ask"]:
            assert q["ask"] in all_text
        for option in q["options"]:
            assert option in all_text
        for analysis in q["analysis"]:
            assert analysis in all_text
    table = doc.tables[0]
    assert len(table.rows) == 8 and len(table.columns) == 10
    for n, q in enumerate(QUESTIONS, start=1):
        col = (n - 1) // 7 * 2
        row = (n - 1) % 7 + 1
        assert table.cell(row, col).text == str(n)
        assert table.cell(row, col + 1).text == q["answer"]
    return doc


def main():
    assert [q["n"] for q in QUESTIONS] == list(range(1, 36))
    assert Counter(q["category"] for q in QUESTIONS) == {
        "图形推理": 10, "定义判断": 10, "类比推理": 5, "逻辑判断": 10
    }
    assert all(q["answer"] in "ABCD" for q in QUESTIONS)
    assert all(q["analysis"] and q["point"] for q in QUESTIONS)
    assert all(q["level"] in {"基础", "中等", "较难"} for q in QUESTIONS)
    assert Counter(q["level"] for q in GRAPHIC) == {"基础": 1, "中等": 6, "较难": 3}
    assert all(len(q["options"]) == (0 if q["n"] <= 7 else 4) for q in QUESTIONS)
    verify_figures()
    assert verify_common()
    doc = verify_docx()
    report = {
        "passed": True,
        "questions": 35,
        "categories": dict(Counter(q["category"] for q in QUESTIONS)),
        "answers": dict(Counter(q["answer"] for q in QUESTIONS)),
        "figure_items": 10,
        "figure_estimated_levels": dict(Counter(q["level"] for q in GRAPHIC)),
        "figure_uniqueness": "通过：1—10题的生成数据、空间约束和分组性质逐题检查",
        "common_cases": 10,
        "common_uniqueness": "通过：31—35题有限状态枚举",
        "word_question_paragraphs": 35,
        "word_answer_paragraphs": 35,
        "word_images": len(doc.inline_shapes),
        "word_sha256": sha256(DOCX),
    }
    (ROOT / "verification.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
