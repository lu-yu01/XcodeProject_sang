from pathlib import Path
from io import BytesIO
import json,hashlib,zipfile,re,collections
from docx import Document
from pypdf import PdfReader
from PIL import Image
from model import Q,verify
ROOT=Path(__file__).resolve().parent;PROJECT=ROOT.parents[1]
FINAL=PROJECT/'模式题/北京风格情境方向图形推理-10题（含答案解析）.docx'
STAGED=ROOT/'revised.docx';SOURCE=ROOT/'source.docx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(n):return json.loads((ROOT/n).read_text(encoding='utf8'))
def norm(t):return re.sub(r'\s+','',t)
def run():
    binding=read('word_bindings.json');assert binding['source_sha256']==sha(SOURCE) and binding['revised_sha256']==sha(STAGED)
    assert verify()['passed']
    with zipfile.ZipFile(SOURCE) as a,zipfile.ZipFile(STAGED) as b:
        assert set(a.namelist())==set(b.namelist())
        for part in binding['preserve_only_identical']:assert a.read(part)==b.read(part),part
        for item in binding['bindings']:
            p=Path(item['source']);assert sha(p)==item['sha256']
            assert hashlib.sha256(b.read(item['word_part'])).hexdigest()==sha(p)
            with Image.open(p) as im:assert abs(item['cx']/item['cy']-im.width/im.height)<.0001
        assert len([n for n in b.namelist() if n.startswith('word/media/')])==20
    old=Document(SOURCE);doc=Document(STAGED)
    assert len(doc.inline_shapes)==20 and len(doc.sections)==len(old.sections)==1
    assert doc.sections[0]._sectPr.xml==old.sections[0]._sectPr.xml
    assert len(doc.paragraphs)==len(old.paragraphs)
    changed=[i for i,(a,b) in enumerate(zip(old.paragraphs,doc.paragraphs)) if a.text!=b.text]
    assert len(changed)==2
    for i in changed:assert doc.paragraphs[i].text in [Q[3]['stem'],Q[6]['stem']]
    texts=[p.text for p in doc.paragraphs]
    for q in Q:
        assert q['stem'] in texts and q['explanation'] in texts
        assert '干扰项辨析：'+q['distractors'] in texts
        assert f'第{q["n"]}题 {q["answer"]}' in texts
    assert '每次转过相同角度' not in ''.join(texts)
    pdf=PdfReader(ROOT/'preview.pdf');assert len(pdf.pages)==10
    for p in pdf.pages:assert len(p.images)==2
    for q in Q:
        n=q['n'];p=pdf.pages[(n-1)//2];a=pdf.pages[5+(n-1)//2]
        assert norm(q['stem']) in norm(p.extract_text()),('stem',n)
        assert norm(f'{n}.') in norm(p.extract_text())
        assert norm(f'第{n}题 {q["answer"]}') in norm(a.extract_text())
        assert norm(q['explanation']) in norm(a.extract_text()),('explanation',n)
        assert norm('干扰项辨析：'+q['distractors']) in norm(a.extract_text()),('distractor',n)
        if n<=5:
            for letter,value in zip('ABCD',q['options']):assert norm(f'{letter}. {value}') in norm(p.extract_text())
    baseline=read('baseline.json');unchanged=[]
    for name,h in baseline.items():
        p=PROJECT/name
        if p.resolve()==FINAL.resolve():continue
        assert sha(p)==h,name;unchanged.append(name)
    assert len(unchanged)==12
    # The untouched first answers are retained even where later picture/condition fixes were needed.
    first1=read('life_first_01_05.json');first2=read('life_first_06_10.json')
    first={q['n']:q['first_option'] for q in first1['questions']}
    first.update({q['n']:q['answer'] for q in first2['answers']})
    assert first=={q['n']:q['answer'] for q in Q}
    close1=read('life_revision_closure_01_05.json')['revision_closure']
    assert close1['summary']['new_delivery_blockers']==[]
    for c in close1['questions']:
        assert c['image_actually_reopened'] and c['option_unique'] and c['blocking_issues_remaining']==[]
        assert c['current_image_sha256'].lower()==sha(ROOT/'assets'/f'question-{c["n"]:02}.png')
        assert c['current_stem']==Q[c['n']-1]['stem']
    for q in first1['questions']:
        if q['n']!=1:assert q['image_sha256'].lower()==sha(ROOT/'assets'/f'question-{q["n"]:02}.png')
    for q in first2['answers']:
        if q['n']!=7:assert q['image_sha256'].lower()==sha(ROOT/'assets'/f'question-{q["n"]:02}.png')
    # Q7 closure has an independent schema; required facts are also checked below.
    close2=read('life_revision_closure_06_10.json')
    c=close2['question']
    assert c['actually_reopened'] and c['closure_answer']=='D' and c['status']=='pass'
    assert c['unique_from_current_image'] and c['previous_issue_resolved']
    assert c['remaining_ambiguities']==[] and c['required_revisions']==[]
    assert c['image_sha256'].lower()==sha(ROOT/'assets'/'question-07.png')
    assert c['stem']==Q[6]['stem']
    for name,h in close2['preserved_files_sha256'].items():assert sha(ROOT/name)==h.lower()
    v=read('visual_review.json')
    assert v['word_sha256']==sha(STAGED) and v['pdf_sha256']==sha(ROOT/'preview.pdf')
    assert len(v['pages'])==10
    for p in v['pages']:
        assert p['actually_opened_at_original_size'] and p['result']=='pass'
        path=Path(p['path'])
        if path.exists():assert sha(path)==p['sha256']
        else:assert any(x['path']==str(path) and x['sha256']==p['sha256'] for x in read('cleanup.json')['deleted'])
    sources=read('source_manifest.json');assert len(sources)==9
    for source in sources:assert sha(ROOT/'sources'/f'{source["name"]}.svg')==source['original_svg_sha256']
    selected=read('selected_assets.json');assert len(selected)==10
    for x in selected:
        assert x['sha256']==sha(Path(x['workspace_asset']))
        assert sha(Path(x['generated_path']))==x['sha256']
    index=PROJECT/'知识库/套卷索引/北京风格情境方向图形推理-10题（含答案解析）.md'
    assert sha(STAGED) in index.read_text(encoding='utf8')
    report={'passed':True,'staged_word_sha256':sha(STAGED),'source_word_sha256':sha(SOURCE),'pdf_sha256':sha(ROOT/'preview.pdf'),
        'questions':10,'answers_unchanged':True,'question_illustrations_replaced':10,'explanation_illustrations_replaced':10,
        'formal_assets':20,'searched_source_vectors':9,'preserved_package_parts':len(binding['preserve_only_identical']),
        'changed_stems':[4,7],'all_figures_match_word':True,'ratios_preserved':True,'final_pdf_pages':10,
        'all_pages_actually_opened':True,'other_word_files_unchanged':unchanged,'first_selections_match_original_answers':10,
        'revision_closures':[1,4,7],'remaining_delivery_blockers':[],
        'limits':'Final artistic question PNGs were independently read for qualitative direction and uniqueness. Numerical model checks validate underlying native scenes, not exact photogrammetric fidelity of generated strokes. No learner difficulty calibration.',
        'audit_fingerprints':{n:sha(ROOT/n) for n in ['life_first_01_05.json','life_first_06_10.json','life_blind_01_05.json','life_blind_06_10.json','life_revision_closure_01_05.json','life_revision_closure_06_10.json','visual_review.json','source_manifest.json','selected_assets.json']}}
    if sha(FINAL)==sha(STAGED):report['final_word_verified']=str(FINAL)
    else:assert sha(FINAL)==sha(SOURCE),'Target changed externally; do not overwrite'
    (ROOT/'final_check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k:report[k] for k in ['passed','staged_word_sha256','final_pdf_pages','formal_assets']},ensure_ascii=False))
if __name__=='__main__':run()
