"""Build the corrected ten-item exercise with the exact requested question form."""
from pathlib import Path
from collections import Counter
import hashlib,json,sys
from docx import Document
from docx.shared import Mm,Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from layout import style,Numbers,keep,page_field,picture
from puzzles import QUESTIONS
import draw

ROOT=Path(__file__).resolve().parent;PROJECT=ROOT.parent.parent
TITLE='立体拼合三黑块专项练习10题'
DOCX=PROJECT/'模式题'/'立体拼合专项练习-三黑块版10题（选项横框版，含答案解析）.docx'
MD=PROJECT/'知识库'/'套卷索引'/'立体拼合专项练习-三黑块版10题（选项横框版，含答案解析）.md'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def answer_table(doc):
    t=doc.add_table(rows=2,cols=10);t.autofit=False;pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','5'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    pr.append(borders)
    mar=OxmlElement('w:tblCellMar')
    for side in ('top','bottom','start','end'):
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
    pr.append(mar)
    for i,q in enumerate(QUESTIONS):
        for r,value in [(0,str(q['n'])),(1,q['answer'])]:
            c=t.cell(r,i);c.width=Inches(.66)
            p=c.paragraphs[0];p.style='Small';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(value).bold=True
            if r==0:
                s=OxmlElement('w:shd');s.set(qn('w:fill'),'E8E8E8');c._tc.get_or_add_tcPr().append(s)

def main():
    sys.stdout.reconfigure(encoding='utf-8')
    assert len(QUESTIONS)==10
    for q in QUESTIONS:
        assert len(q['pieces'])==2 and len(q['options'])==4
        assert len(q['target']['black'])==3
        assert all(len(s['black'])==1 for s in q['pieces']+q['options'])
    draw.draw_all(QUESTIONS)
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(18)
    sec.footer_distance=Mm(8)
    doc.core_properties.title=TITLE;doc.core_properties.subject='完整目标与两个部分 选择第三部分 三黑块位置保持';doc.core_properties.author='教研资料'
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    style(doc.styles['Normal'],11)
    style(doc.styles['Title'],18,True,0,5,1.1,'黑体')
    for name,opts in {
        'Question':dict(size=11,before=5,after=3,line=1.15),
        'Figure':dict(size=1,before=2,after=3,line=1),
        'Small':dict(size=9.5,before=0,after=3,line=1.1),
        'AnswerTitle':dict(size=11,bold=True,before=6,after=3,east='黑体'),
        'Point':dict(size=10,bold=True,before=0,after=3),
        'Analysis':dict(size=10.5,before=0,after=3,line=1.15)
    }.items():
        s=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);s.base_style=doc.styles['Normal'];style(s,**opts)
    for name in ('AnswerTitle','Point'):doc.styles[name].paragraph_format.keep_with_next=True
    style(doc.styles['Footer'],8,False,0,0,1)
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('第 ');page_field(p);p.add_run(' 页')
    p=doc.add_paragraph(TITLE,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('每题给出完整立体和部件①、②，从A—D中选择第三部分，使三件共同拼成目标。目标共3个黑块，每件各1个黑块；白块均为白色，黑色为固定的材料标记。','Small')
    doc.add_paragraph('允许部件整体旋转、平移；不得镜像、拆分、变形、重叠或留下空缺。完整立体左、右两幅分别为正面、背面；必要块数列在图片外。','Small')
    doc.add_paragraph('姓名：________________　日期：________________　得分：________','Small')
    nums=Numbers(doc);qid=nums.make();aid=nums.make()
    for i,q in enumerate(QUESTIONS):
        p=doc.add_paragraph(q['stem'],'Question');nums.apply(p,qid)
        if i and i%2==0:p.paragraph_format.page_break_before=True
        count=doc.add_paragraph(f'体块数：目标{len(q["target"]["cells"])}块；①{len(q["pieces"][0]["cells"])}块；②{len(q["pieces"][1]["cells"])}块；A—D各{len(q["options"][0]["cells"])}块。','Small')
        fig=picture(doc,draw.ASSETS/f'q{q["n"]:02}.png',f'第{q["n"]}题 完整目标三黑块 部件一二各一黑块 四候选各一黑块',6.75,3.5)
        keep([p,count,fig])
    p=doc.add_paragraph('参考答案与解析','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before=True
    answer_table(doc)
    doc.add_paragraph('下方分层图用①、②及答案字母区分三个部件，黑格对应目标黑块；各层按同一朝向自底向上叠放，空格在目标之外。','Small')
    for i,q in enumerate(QUESTIONS):
        block=[]
        p=doc.add_paragraph('答案：'+q['answer'],'AnswerTitle');nums.apply(p,aid);block.append(p)
        if i and i%2==0:p.paragraph_format.page_break_before=True
        block.append(doc.add_paragraph('考点：'+q['point'],'Point'))
        for j,s in enumerate(q['analysis']):block.append(doc.add_paragraph(('解析：' if j==0 else '')+s,'Analysis'))
        block.append(picture(doc,draw.ASSETS/f'q{q["n"]:02}_solution.png',f'第{q["n"]}题三部件与三个黑块对应的分层构造',6.6,1.65))
        keep(block)
    DOCX.parent.mkdir(exist_ok=True);doc.save(DOCX)
    info={'docx':str(DOCX),'word_sha256':sha(DOCX),'questions':10,'embedded_images':len(doc.inline_shapes),
          'answers':dict(Counter(q['answer'] for q in QUESTIONS)),'shape_form':'complete target + 2 given parts + choose third','black_counts':'target3 / given1+1 / option1'}
    assert info['embedded_images']==20
    (ROOT/'build_info.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
    MD.write_text(f'''# 立体拼合三黑块版10题 索引与检查记录

- Word成品：[三黑块版10题](../../模式题/立体拼合专项练习-三黑块版10题（选项横框版，含答案解析）.docx)。
- 工作目录：[assembly_black10_options_20261025](../../.work/assembly_black10_options_20261025/)。
- 题量：10题全部为完整目标加两个已给部分，选择第三部分；目标3黑块，已给各1黑块，每个候选各1黑块。
- 当前状态：已生成，待数学、颜色对应、逐图与成品逐页验收。
- Word版本指纹：`{sha(DOCX)}`。
''',encoding='utf-8')
    print(json.dumps(info,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
